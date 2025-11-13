"""
Metadata Storage Backend
Handles persistence of metadata to JSON files and provides CRUD operations
"""
import json
import os
from typing import List, Dict, Optional
from pathlib import Path
from metadata_models import (
    PipelineMetadata, SourceMetadata, TransformationMetadata,
    DataQualityRule, TargetMetadata, ColumnMetadata
)


class MetadataStore:
    """Storage layer for pipeline metadata"""
    
    def __init__(self, storage_path: str = "./metadata_store"):
        self.storage_path = Path(storage_path)
        self.storage_path.mkdir(exist_ok=True, parents=True)
        self.pipelines_path = self.storage_path / "pipelines"
        self.pipelines_path.mkdir(exist_ok=True)
        
    def save_pipeline(self, pipeline: PipelineMetadata) -> bool:
        """Save pipeline metadata to storage"""
        try:
            file_path = self.pipelines_path / f"{pipeline.pipeline_id}.json"
            with open(file_path, 'w') as f:
                json.dump(pipeline.to_dict(), f, indent=2)
            return True
        except Exception as e:
            print(f"Error saving pipeline: {e}")
            return False
    
    def load_pipeline(self, pipeline_id: str) -> Optional[Dict]:
        """Load pipeline metadata from storage"""
        try:
            file_path = self.pipelines_path / f"{pipeline_id}.json"
            if not file_path.exists():
                return None
            with open(file_path, 'r') as f:
                return json.load(f)
        except Exception as e:
            print(f"Error loading pipeline: {e}")
            return None
    
    def list_pipelines(self) -> List[str]:
        """List all pipeline IDs"""
        return [f.stem for f in self.pipelines_path.glob("*.json")]
    
    def delete_pipeline(self, pipeline_id: str) -> bool:
        """Delete pipeline metadata"""
        try:
            file_path = self.pipelines_path / f"{pipeline_id}.json"
            if file_path.exists():
                file_path.unlink()
                return True
            return False
        except Exception as e:
            print(f"Error deleting pipeline: {e}")
            return False
    
    def search_pipelines(self, name: str = None, tags: List[str] = None) -> List[Dict]:
        """Search pipelines by name or tags"""
        results = []
        for pipeline_id in self.list_pipelines():
            pipeline = self.load_pipeline(pipeline_id)
            if pipeline:
                match = True
                if name and name.lower() not in pipeline['name'].lower():
                    match = False
                if tags and not any(tag in pipeline.get('tags', []) for tag in tags):
                    match = False
                if match:
                    results.append(pipeline)
        return results
    
    def export_to_excel(self, pipeline_id: str, output_path: str) -> bool:
        """Export pipeline metadata to Excel format (for compatibility)"""
        try:
            import pandas as pd
            from openpyxl import Workbook
            from openpyxl.utils.dataframe import dataframe_to_rows
            
            pipeline = self.load_pipeline(pipeline_id)
            if not pipeline:
                return False
            
            wb = Workbook()
            wb.remove(wb.active)
            
            # Sources sheet
            if pipeline['sources']:
                ws_sources = wb.create_sheet("Sources")
                sources_data = []
                for src in pipeline['sources']:
                    for col in src['columns']:
                        sources_data.append({
                            'Source ID': src['source_id'],
                            'Source Name': src['name'],
                            'Location': src['location'],
                            'Format': src['format'],
                            'Column Name': col['name'],
                            'Data Type': col['data_type'],
                            'Nullable': col['nullable'],
                            'Description': col['description']
                        })
                df = pd.DataFrame(sources_data)
                for r in dataframe_to_rows(df, index=False, header=True):
                    ws_sources.append(r)
            
            # Transformations sheet
            if pipeline['transformations']:
                ws_trans = wb.create_sheet("Transformations")
                trans_data = [{
                    'Transformation ID': t['transformation_id'],
                    'Name': t['name'],
                    'Type': t['transformation_type'],
                    'Source References': ', '.join(t['source_refs']),
                    'Logic': t['logic'],
                    'Description': t['description']
                } for t in pipeline['transformations']]
                df = pd.DataFrame(trans_data)
                for r in dataframe_to_rows(df, index=False, header=True):
                    ws_trans.append(r)
            
            # Data Quality Rules sheet
            if pipeline['data_quality_rules']:
                ws_dq = wb.create_sheet("Data Quality Rules")
                dq_data = [{
                    'Rule ID': r['rule_id'],
                    'Name': r['name'],
                    'Type': r['rule_type'],
                    'Target Columns': ', '.join(r['target_columns']),
                    'Condition': r['condition'],
                    'Severity': r['severity'],
                    'Enabled': r['enabled']
                } for r in pipeline['data_quality_rules']]
                df = pd.DataFrame(dq_data)
                for r in dataframe_to_rows(df, index=False, header=True):
                    ws_dq.append(r)
            
            # Targets sheet
            if pipeline['targets']:
                ws_targets = wb.create_sheet("Targets")
                targets_data = [{
                    'Target ID': t['target_id'],
                    'Name': t['name'],
                    'Location': t['location'],
                    'Format': t['format'],
                    'Write Mode': t['write_mode'],
                    'Source Transformation': t['source_transformation_ref'],
                    'Description': t['description']
                } for t in pipeline['targets']]
                df = pd.DataFrame(targets_data)
                for r in dataframe_to_rows(df, index=False, header=True):
                    ws_targets.append(r)
            
            wb.save(output_path)
            return True
        except Exception as e:
            print(f"Error exporting to Excel: {e}")
            return False
